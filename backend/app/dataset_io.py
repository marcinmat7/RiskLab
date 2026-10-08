import csv
import io
import json
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pyarrow.feather as feather
import pyarrow.parquet as parquet
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook

SUPPORTED_EXTENSIONS = {
    ".csv": "Delimited text",
    ".tsv": "Delimited text",
    ".txt": "Delimited text",
    ".psv": "Delimited text",
    ".parquet": "Parquet",
    ".feather": "Feather",
    ".xlsx": "Excel",
    ".jsonl": "JSON Lines",
    ".ndjson": "JSON Lines",
}

DELIMITERS = ",;\t|"


def supported_extensions_label() -> str:
    return ", ".join(sorted(SUPPORTED_EXTENSIONS))


def _stringify(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    return str(value)


def _validate_columns(columns: list[str]) -> list[str]:
    cleaned = [str(column).strip() for column in columns]
    if not cleaned or all(not column for column in cleaned):
        raise HTTPException(status_code=400, detail="The dataset must contain column names.")
    if any(not column for column in cleaned):
        raise HTTPException(status_code=400, detail="Empty column names are not supported.")
    if len(set(cleaned)) != len(cleaned):
        raise HTTPException(status_code=400, detail="Duplicate column names are not supported.")
    if len(cleaned) < 2:
        raise HTTPException(status_code=400, detail="The dataset must contain at least two columns.")
    return cleaned


async def _read_bytes(file: UploadFile) -> bytes:
    await file.seek(0)
    payload = await file.read()
    if not payload:
        raise HTTPException(status_code=400, detail="The uploaded file is empty.")
    return payload


def _read_delimited(payload: bytes) -> tuple[list[str], list[dict[str, str]], str]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="Delimited text files must be UTF-8 encoded.") from exc

    sample = text[:8192]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=DELIMITERS)
    except csv.Error:
        dialect = csv.excel

    reader = csv.DictReader(io.StringIO(text), dialect=dialect)
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="The dataset must contain a header row.")

    columns = _validate_columns(list(reader.fieldnames))
    rows = [
        {column: _stringify(raw.get(column)) for column in columns}
        for raw in reader
    ]
    return columns, rows, dialect.delimiter


def _read_parquet(payload: bytes) -> tuple[list[str], list[dict[str, str]], str]:
    try:
        table = parquet.read_table(io.BytesIO(payload))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read Parquet file: {exc}") from exc
    columns = _validate_columns(table.column_names)
    rows = [{column: _stringify(row.get(column)) for column in columns} for row in table.to_pylist()]
    return columns, rows, ""


def _read_feather(payload: bytes) -> tuple[list[str], list[dict[str, str]], str]:
    try:
        table = feather.read_table(io.BytesIO(payload))
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read Feather file: {exc}") from exc
    columns = _validate_columns(table.column_names)
    rows = [{column: _stringify(row.get(column)) for column in columns} for row in table.to_pylist()]
    return columns, rows, ""


def _read_xlsx(payload: bytes) -> tuple[list[str], list[dict[str, str]], str]:
    try:
        workbook = load_workbook(io.BytesIO(payload), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read Excel file: {exc}") from exc

    worksheet = workbook.active
    iterator = worksheet.iter_rows(values_only=True)
    header = next(iterator, None)
    if header is None:
        raise HTTPException(status_code=400, detail="The Excel workbook is empty.")

    columns = _validate_columns([_stringify(value) for value in header])
    rows = []
    for values in iterator:
        padded = list(values) + [None] * max(0, len(columns) - len(values))
        rows.append({column: _stringify(padded[index]) for index, column in enumerate(columns)})
    workbook.close()
    return columns, rows, ""


def _read_json_lines(payload: bytes) -> tuple[list[str], list[dict[str, str]], str]:
    try:
        text = payload.decode("utf-8-sig")
    except UnicodeDecodeError as exc:
        raise HTTPException(status_code=400, detail="JSON Lines files must be UTF-8 encoded.") from exc

    objects: list[dict[str, Any]] = []
    columns: list[str] = []
    seen: set[str] = set()

    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        if not raw_line.strip():
            continue
        try:
            value = json.loads(raw_line)
        except json.JSONDecodeError as exc:
            raise HTTPException(status_code=400, detail=f"Invalid JSON on line {line_number}.") from exc
        if not isinstance(value, dict):
            raise HTTPException(status_code=400, detail=f"JSON Lines row {line_number} must be an object.")
        objects.append(value)
        for key in value:
            column = str(key).strip()
            if column and column not in seen:
                seen.add(column)
                columns.append(column)

    columns = _validate_columns(columns)
    rows = [{column: _stringify(obj.get(column)) for column in columns} for obj in objects]
    return columns, rows, ""


async def load_tabular_file(file: UploadFile) -> dict[str, Any]:
    filename = file.filename or "dataset"
    extension = Path(filename).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file type '{extension or '(none)'}'. Supported: {supported_extensions_label()}.",
        )

    payload = await _read_bytes(file)

    if extension in {".csv", ".tsv", ".txt", ".psv"}:
        columns, rows, delimiter = _read_delimited(payload)
    elif extension == ".parquet":
        columns, rows, delimiter = _read_parquet(payload)
    elif extension == ".feather":
        columns, rows, delimiter = _read_feather(payload)
    elif extension == ".xlsx":
        columns, rows, delimiter = _read_xlsx(payload)
    else:
        columns, rows, delimiter = _read_json_lines(payload)

    if not rows:
        raise HTTPException(status_code=400, detail="The dataset contains column names but no data rows.")

    return {
        "filename": filename,
        "extension": extension,
        "format": SUPPORTED_EXTENSIONS[extension],
        "delimiter": delimiter,
        "columns": columns,
        "rows": rows,
        "file_size_bytes": len(payload),
    }
